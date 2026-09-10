"""Bug 4: Page propagation across ActionGroup and Workflow."""

from __future__ import annotations

from models import ActionGroup, ActionRecord
from segmentation import WorkflowSegmenter


def _action(action_id: str, page: int, order: int) -> ActionRecord:
    return ActionRecord(
        action_id=action_id,
        action="inspect",
        object="crack",
        sentence=f"Inspect on page {page}",
        page=page,
        document_id="doc1",
        source_name="manual.docx",
        chunk_id=f"chunk_{page}",
        order_index=order,
    )


def test_action_group_pages_span_multiple_pages():
    actions = [_action("a1", 2, 0), _action("a2", 2, 1), _action("a3", 3, 2)]
    group = ActionGroup(
        document_id="doc1",
        source_name="manual.docx",
        actions=actions,
    )
    segmenter = WorkflowSegmenter()
    groups = segmenter._split_into_groups(actions, [], [], {})  # type: ignore[arg-type]
    assert len(groups) == 1
    assert groups[0].pages == [2, 3]
    assert groups[0].page_start == 2
    assert groups[0].page_end == 3


def test_docx_virtual_pages_increment_on_major_sections(tmp_path):
    from parser import DocumentParserService

    docx_path = tmp_path / "manual.docx"
    from docx import Document

    document = Document()
    document.add_heading("4.3  Fuselage Skin Crack — Field Repair Procedure", level=1)
    document.add_paragraph("1. Inspect the crack.")
    document.add_heading("4.5  Engine Mount Bracket Weld Repair", level=1)
    document.add_paragraph("1. Inspect the crack region.")
    document.save(str(docx_path))

    parsed = DocumentParserService().parse(docx_path)
    pages = {u.text[:20]: u.page for u in parsed.units if u.unit_type != "heading"}
    inspect_pages = [u.page for u in parsed.units if "Inspect the crack" in u.text]
    assert 1 in inspect_pages
    assert 2 in inspect_pages
